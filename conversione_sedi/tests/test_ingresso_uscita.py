import openpyxl
import pandas as pd
import pytest

from conversione_sedi.cli import CODICE_ERRORE, CODICE_OK, main
from conversione_sedi.ingresso_uscita import (
    REGOLE_PREDEFINITE,
    TRASFORMAZIONI_PREDEFINITE,
    ErroreDati,
    converti_file,
    elabora,
    leggi_sedi,
    percorso_output_predefinito,
    scrivi_output,
    trova_colonna_sede,
)
from conversione_sedi.nucleo import (
    ESITO_DA_VERIFICARE, ESITO_NESSUNA_CORRISPONDENZA, ESITO_OK, ESITO_VUOTA,
    ErroreConfigurazione, carica_configurazione,
)

SEDI = ["ITE-OT0B-1E-01-B1-BM", "ITE-OT0B-00", "", "ITE-OT0B-0E-SA", "ITE-OT0B-1E-MV-I1"]
DESCRIZIONI = ["Batteria", "Sconosciuta", "Vuota", "Servizi", "01"]


@pytest.fixture
def conf():
    return carica_configurazione(REGOLE_PREDEFINITE, TRASFORMAZIONI_PREDEFINITE)


def crea_excel(percorso, colonne):
    pd.DataFrame(colonne).to_excel(percorso, index=False)
    return percorso


# ---------------------------------------------------------------------------
# Configurazione predefinita
# ---------------------------------------------------------------------------
def test_config_predefinita_valida(conf):
    assert conf.num_regole == 11
    assert conf.num_trasformazioni == 6
    assert conf.avvisi == []


# ---------------------------------------------------------------------------
# Individuazione della colonna
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("nome", ["Sede tecnica", "Sede_tecnica", " SEDE TECNICA ", "sede_Tecnica"])
def test_colonna_predefinita(nome):
    df = pd.DataFrame({"Altro": ["x"], nome: ["y"]})
    assert trova_colonna_sede(df) == nome


def test_colonna_esplicita():
    df = pd.DataFrame({"FL": ["x"], "Sede tecnica": ["y"]})
    assert trova_colonna_sede(df, "fl") == "FL"


def test_colonna_mancante():
    with pytest.raises(ErroreDati, match="non trovata.*\"Altro\""):
        trova_colonna_sede(pd.DataFrame({"Altro": ["x"]}))


def test_colonna_ambigua():
    df = pd.DataFrame([["a", "b"]], columns=["Sede tecnica", "Sede_tecnica"])
    with pytest.raises(ErroreDati, match="più volte"):
        trova_colonna_sede(df)


# ---------------------------------------------------------------------------
# Lettura
# ---------------------------------------------------------------------------
def test_leggi_excel_come_testo(tmp_path):
    p = tmp_path / "sedi.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Sede tecnica", "Descrizione", "Codice"])
    ws.append(["ITE-OT0B-0E", "NA", "01"])        # "01" come testo
    ws.append(["ITE-OT0B-0E-SA", "nn", None])
    wb.save(p)
    df = leggi_sedi(p)
    assert list(df["Descrizione"]) == ["NA", "nn"]     # nessuna conversione in NaN
    assert list(df["Codice"]) == ["01", ""]


def test_leggi_excel_foglio(tmp_path):
    p = tmp_path / "sedi.xlsx"
    with pd.ExcelWriter(p) as w:
        pd.DataFrame({"X": ["1"]}).to_excel(w, sheet_name="Primo", index=False)
        pd.DataFrame({"Sede tecnica": ["ITE-OT0B-0E"]}).to_excel(w, sheet_name="Sedi", index=False)
    assert list(leggi_sedi(p, "Sedi").columns) == ["Sede tecnica"]
    with pytest.raises(ErroreDati, match="Inesistente"):
        leggi_sedi(p, "Inesistente")


def test_leggi_csv_cp1252_virgola(tmp_path):
    p = tmp_path / "sedi.csv"
    p.write_text("Sede tecnica,Descrizione\nITE-OT0B-0E,Perché\nITE-OT0B-0E-N0,NA\n", encoding="cp1252")
    df = leggi_sedi(p)
    assert list(df["Descrizione"]) == ["Perché", "NA"]


def test_leggi_file_mancante(tmp_path):
    with pytest.raises(ErroreDati, match="File non trovato"):
        leggi_sedi(tmp_path / "manca.xlsx")


def test_leggi_csv_vuoto(tmp_path):
    p = tmp_path / "vuoto.csv"
    p.write_text("", encoding="utf-8")
    with pytest.raises(ErroreDati, match="vuoto"):
        leggi_sedi(p)


def test_leggi_formato_non_supportato(tmp_path):
    p = tmp_path / "sedi.txt"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(ErroreDati, match="non supportato"):
        leggi_sedi(p)


def test_leggi_excel_danneggiato(tmp_path):
    p = tmp_path / "rotto.xlsx"
    p.write_text("non sono un file Excel", encoding="utf-8")
    with pytest.raises(ErroreDati, match="Impossibile leggere"):
        leggi_sedi(p)


# ---------------------------------------------------------------------------
# Scrittura
# ---------------------------------------------------------------------------
def test_percorso_output_predefinito(tmp_path):
    assert percorso_output_predefinito(tmp_path / "a.xlsx").name == "a_convertito.xlsx"
    assert percorso_output_predefinito(tmp_path / "a.XLSM").name == "a_convertito.xlsx"
    assert percorso_output_predefinito(tmp_path / "a.csv").name == "a_convertito.csv"


def test_scrivi_excel_celle_testo(tmp_path):
    df = pd.DataFrame({"A": ["++E-++++-1A", "01", "=SOMMA(1;2)", None]})
    p = scrivi_output(df, tmp_path / "out.xlsx")
    ws = openpyxl.load_workbook(p).active
    valori = [ws.cell(r, 1) for r in range(2, 5)]
    assert [c.value for c in valori] == ["++E-++++-1A", "01", "=SOMMA(1;2)"]
    assert all(c.data_type == "s" and c.number_format == "@" for c in valori)


def test_scrivi_csv(tmp_path):
    df = pd.DataFrame({"A": ["01", None], "B": ["nn", "x"]})
    p = scrivi_output(df, tmp_path / "out.csv")
    assert p.read_text(encoding="utf-8-sig") == "A;B\n01;nn\n;x\n"


def test_scrivi_formato_non_supportato(tmp_path):
    with pytest.raises(ErroreDati, match="non supportato"):
        scrivi_output(pd.DataFrame({"A": ["x"]}), tmp_path / "out.json")


def test_scrivi_cartella_inesistente(tmp_path):
    with pytest.raises(ErroreDati, match="inesistente"):
        scrivi_output(pd.DataFrame({"A": ["x"]}), tmp_path / "manca" / "out.csv")


# ---------------------------------------------------------------------------
# Elaborazione completa
# ---------------------------------------------------------------------------
def test_elabora_avvisa_colonne_sovrascritte(conf):
    df = pd.DataFrame({"Sede tecnica": ["ITE-OT0B-0E"], "Esito": ["vecchio"]})
    esito = elabora(df, conf)
    assert esito.dati.loc[0, "Esito"] == ESITO_OK
    assert any("Esito" in a for a in esito.avvisi)


def test_converti_file_excel(tmp_path):
    ingresso = crea_excel(tmp_path / "sedi.xlsx", {"Sede tecnica": SEDI, "Descrizione": DESCRIZIONI})
    esito = converti_file(ingresso)
    assert esito.percorso_output == tmp_path / "sedi_convertito.xlsx"
    out = leggi_sedi(esito.percorso_output)
    assert list(out.columns) == ["Sede tecnica", "Descrizione", "Codice_Out", "Nuova_Sede_Tecnica", "Esito"]
    assert list(out["Descrizione"]) == DESCRIZIONI
    assert list(out["Nuova_Sede_Tecnica"]) == ["ITE-OT0B-1E-01-C1-R1", "", "", "ITE-OT0B-1A-S1",
                                               "ITE-OT0B-1E-01-MV-RM"]
    assert list(out["Esito"]) == [ESITO_OK, ESITO_NESSUNA_CORRISPONDENZA, ESITO_VUOTA, ESITO_OK, ESITO_OK]
    assert esito.riepilogo == {ESITO_OK: 3, ESITO_DA_VERIFICARE: 0,
                               ESITO_NESSUNA_CORRISPONDENZA: 1, ESITO_VUOTA: 1}


def test_converti_file_csv_verso_excel(tmp_path):
    ingresso = tmp_path / "sedi.csv"
    ingresso.write_text("FL;Note\nITE-OT0B-1E-01-B1-BS;a\n", encoding="utf-8")
    esito = converti_file(ingresso, tmp_path / "risultato.xlsx", colonna="FL")
    assert leggi_sedi(esito.percorso_output).loc[0, "Nuova_Sede_Tecnica"] == "ITE-OT0B-1E-01-C1-BS"


def test_converti_file_regole_personalizzate(tmp_path):
    regole = tmp_path / "r.csv"
    regole.write_text("Regola_In;Codice_Out\n++E-++++-+E-MV-I+;++E-++++-+E-nn-MV-RM\n", encoding="utf-8")
    trasf = tmp_path / "t.csv"
    trasf.write_text("Segmento;Valore_Atteso;Espressione\n1;*;{S1}\nx;;\n", encoding="utf-8")
    ingresso = tmp_path / "sedi.csv"
    ingresso.write_text("Sede_tecnica\nITE-OT0B-1E-MV-I1\n", encoding="utf-8")
    esito = converti_file(ingresso, percorso_regole=regole, percorso_trasformazioni=trasf)
    assert esito.dati.loc[0, "Esito"] == ESITO_DA_VERIFICARE
    assert esito.dati.loc[0, "Nuova_Sede_Tecnica"].startswith("DA VERIFICARE: ")
    assert any("righe ignorate" in a for a in esito.avvisi)


def test_converti_file_output_uguale_input(tmp_path):
    ingresso = crea_excel(tmp_path / "sedi.xlsx", {"Sede tecnica": SEDI})
    with pytest.raises(ErroreDati, match="coincidere"):
        converti_file(ingresso, ingresso)


def test_converti_file_errore_configurazione(tmp_path):
    ingresso = crea_excel(tmp_path / "sedi.xlsx", {"Sede tecnica": SEDI})
    with pytest.raises(ErroreConfigurazione) as exc:
        converti_file(ingresso, percorso_regole=tmp_path / "a.csv",
                      percorso_trasformazioni=tmp_path / "b.csv")
    assert len(exc.value.errori) == 2


# ---------------------------------------------------------------------------
# Riga di comando
# ---------------------------------------------------------------------------
def test_cli_ok(tmp_path, capsys):
    ingresso = crea_excel(tmp_path / "sedi.xlsx", {"Sede tecnica": SEDI})
    assert main([str(ingresso), "-o", str(tmp_path / "out.csv")]) == CODICE_OK
    stampa = capsys.readouterr().out
    assert "Elaborate 5 righe" in stampa and "NESSUNA_CORRISPONDENZA" in stampa
    assert (tmp_path / "out.csv").is_file()


def test_cli_errore_dati(tmp_path, capsys):
    ingresso = crea_excel(tmp_path / "sedi.xlsx", {"Altro": SEDI})
    assert main([str(ingresso)]) == CODICE_ERRORE
    errori = capsys.readouterr().err
    assert errori.startswith("ERRORE:") and "Traceback" not in errori


def test_cli_errore_configurazione(tmp_path, capsys):
    ingresso = crea_excel(tmp_path / "sedi.xlsx", {"Sede tecnica": SEDI})
    regole = tmp_path / "r.csv"
    regole.write_text("Regola;Codice_Out\nA;B\n", encoding="utf-8")
    assert main([str(ingresso), "--regole", str(regole)]) == CODICE_ERRORE
    errori = capsys.readouterr().err
    assert "ERRORE nei file delle regole" in errori and "Regola_In" in errori


def test_cli_avvisi(tmp_path, capsys):
    ingresso = crea_excel(tmp_path / "sedi.xlsx", {"Sede tecnica": SEDI})
    trasf = tmp_path / "t.csv"
    trasf.write_text("Segmento;Valore_Atteso;Espressione\n1;*;{S1}\n0;*;X\n", encoding="utf-8")
    assert main([str(ingresso), "--trasformazioni", str(trasf)]) == CODICE_OK
    assert "Avvisi:" in capsys.readouterr().err
