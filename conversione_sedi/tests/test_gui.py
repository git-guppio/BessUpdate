"""Test dell'interfaccia grafica senza finestre visibili (Qt offscreen)."""
import os

import pandas as pd
import pytest

pytest.importorskip("PySide6")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook                                     # noqa: E402
from PySide6.QtCore import Qt                                          # noqa: E402
from PySide6.QtWidgets import QApplication                             # noqa: E402

from conversione_sedi.esportazione import esporta_excel, testo_tabulato as tabulato  # noqa: E402
from conversione_sedi.gui.attesa import esegui_con_attesa              # noqa: E402
from conversione_sedi.gui.finestra import Finestra                     # noqa: E402
from conversione_sedi.gui.stile import BLU, GIALLO, GRIGIO, ROSSO, Stile, sfondo_esito  # noqa: E402
from conversione_sedi.nuova_codifica import valida_new                 # noqa: E402
from conversione_sedi.progetto import Progetto, Scheda                 # noqa: E402
from conftest import REGOLE_TEST, TRASFORMAZIONI_TEST  # noqa: E402

NEW = [("ITE-OT0B", "Impianto"), ("ITE-OT0B-1E", "Area 1"), ("ITE-OT0B-1E-01", "Linea 1"),
       ("ITE-OT0B-1E-01-C1", "Cluster"), ("ITE-OT0B-1E-01-C1-R1", "Rack 1")]
ACTUAL = [("ITE-OT0B", "Impianto"), ("ITE-OT0B-1E", "Area 1"), ("ITE-OT0B-1E-01", "Linea 1"),
          ("ITE-OT0B-1E-01-B1", "Batteria"), ("ITE-OT0B-1E-01-B1-BM", "Modulo")]


def tabella(righe):
    return pd.DataFrame(righe, columns=["Sede tecnica", "Descrizione"])


@pytest.fixture(scope="session")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def impostazioni(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))


@pytest.fixture
def finestra(app, tmp_path):
    p = Progetto(REGOLE_TEST, TRASFORMAZIONI_TEST)
    p.imposta_new(valida_new(tabella(NEW)))
    p.imposta_actual(tabella(ACTUAL))
    p.salva(tmp_path / "prova.bessprj")
    f = Finestra(p)
    f.errori = []
    f.errore = lambda e: f.errori.append(e)                 # niente finestre modali nei test
    yield f
    f.close()


def pagina(f, scheda):
    i = [pg.scheda for pg in f.pagine].index(scheda)
    f.passi.setCurrentRow(i)
    return f.pagine[i]


# ---------------------------------------------------------------------------
# Stile ed esportazione (senza Qt)
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("valore, atteso", [
    ("OK", None), ("Modifica", None), ("Conflitto", ROSSO), ("Parent NOK (manca X)", ROSSO),
    ("Valore duplicato | Parent NOK (manca X)", ROSSO), ("Mancante in SAP", ROSSO),
    ("Da decidere", GIALLO), ('Descrizione diversa (New: "x")', GIALLO), ("DA_VERIFICARE", GIALLO),
    ("Esclusa (livello 1-2)", GRIGIO), ("NESSUNA_CORRISPONDENZA", GRIGIO), (None, None),
])
def test_sfondo_esito(valore, atteso):
    assert sfondo_esito(valore) == atteso


def test_esporta_excel_con_colori(tmp_path):
    df = pd.DataFrame({"Sede": ["++E-01", "=1+1"], "Esito": ["OK", "Conflitto"]})
    stile = lambda r, c: Stile(BLU, ROSSO if r == 1 else None, r == 0)  # noqa: E731
    f = esporta_excel(df, tmp_path / "out", stile=stile)
    ws = load_workbook(f).active
    assert f.name == "out.xlsx"
    assert [c.value for c in ws[2]] == ["++E-01", "OK"] and ws["A3"].value == "=1+1"
    assert ws["A3"].data_type == "s"                                    # mai formula
    assert ws["A2"].font.strike and ws["A3"].fill.start_color.rgb == "FFF8D7DA"


def test_testo_tabulato():
    df = pd.DataFrame({"A": ["x\ty", None], "B": ["1", "2"]})
    assert tabulato(df) == "A\tB\r\nx y\t1\r\n\t2"


# ---------------------------------------------------------------------------
# Finestra
# ---------------------------------------------------------------------------
def test_barra_laterale(finestra):
    testi = [finestra.passi.item(i).text() for i in range(finestra.passi.count())]
    assert testi[0] == "1  ✓  New"
    assert testi[6] == "7  🔒  No_Wo"                   # riconciliazione incompleta
    assert testi[7] == "8  🔒  SAP script"
    assert testi[8] == "9  ○  Check"
    assert finestra.windowTitle() == "BESS Update — prova.bessprj — ITE-OT0B"


def test_tutte_le_schede_si_aprono(finestra):
    for i, pg in enumerate(finestra.pagine):
        finestra.passi.setCurrentRow(i)
        pg.aggiorna()
    assert finestra.errori == []
    assert pagina(finestra, Scheda.NEW).pannello.modello.rowCount() == 5
    assert pagina(finestra, Scheda.FL_DA_CREARE).pannello.modello.rowCount() == 5
    assert pagina(finestra, Scheda.NO_WO).avviso.text().startswith("Scheda bloccata")


def test_decisioni_in_check_actual(finestra):
    pg = pagina(finestra, Scheda.CHECK_ACTUAL)
    m = pg.pannello.modello
    riga = list(m.df["Sede tecnica"]).index("ITE-OT0B-1E-01-B1")
    assert m.valore(riga, "Operazione") == "Da decidere"

    # menu contestuale: Da rimuovere
    azioni = {a.text(): a for a in pg.azioni([riga], "Operazione") if a is not None}
    azioni["Da rimuovere"].trigger()
    riga = list(m.df["Sede tecnica"]).index("ITE-OT0B-1E-01-B1")
    assert m.valore(riga, "Operazione") == "Da rimuovere"
    assert m.stile(riga, "Sede tecnica").testo == BLU                   # decisione manuale in blu
    assert "Riconciliazione completa" in pg.pannello.info.text()
    assert finestra.passi.item(6).text() == "7  ○  No_Wo"               # sbloccata

    # modifica in cella del nuovo codice
    m.setData(m.index(riga, list(m.df.columns).index("Nuova sede tecnica")), "ite-ot0b-1e-01-c1")
    riga = list(m.df["Sede tecnica"]).index("ITE-OT0B-1E-01-B1")
    assert m.valore(riga, "Operazione") == "Modifica"
    assert finestra.progetto.decisioni["ITE-OT0B-1E-01-B1"].codice == "ITE-OT0B-1E-01-C1"


def test_modifica_manuale_new_e_ricalcolo(finestra):
    pg = pagina(finestra, Scheda.NEW)
    m = pg.pannello.modello
    m.setData(m.index(2, 1), "Linea 1 nuova")
    assert m.valore(2, "Descrizione") == "Linea 1 nuova" and m.valore(2, "Origine") == "MODIFICATO"
    assert m.stile(2, "Descrizione").testo == BLU and m.stile(2, "Sede tecnica").testo is None
    assert m.data(m.index(2, 1), Qt.ToolTipRole) == "Valore importato: Linea 1"
    pg_actual = pagina(finestra, Scheda.CHECK_ACTUAL).pannello.modello
    riga = list(pg_actual.df["Sede tecnica"]).index("ITE-OT0B-1E-01")
    assert pg_actual.valore(riga, "Operazione") == "Modifica descrizione"


def test_errore_di_modifica_mostrato(finestra):
    m = pagina(finestra, Scheda.NEW).pannello.modello
    m.setData(m.index(2, 0), "USE-USS5-1E")                            # altro impianto
    assert "non appartiene all'impianto" in str(finestra.errori[0])


def test_rimozione_e_ripristino_righe(finestra):
    pg = pagina(finestra, Scheda.NEW)
    azioni = {a.text(): a for a in pg.azioni([4], None) if a is not None}
    azioni["Rimuovi righe selezionate"].trigger()
    m = pg.pannello.modello
    assert m.valore(4, "Origine") == "RIMOSSO" and m.data(m.index(4, 0), Qt.FontRole).strikeOut()
    azioni = {a.text(): a for a in pg.azioni([4], None) if a is not None}
    assert not azioni["Rimuovi righe selezionate"].isEnabled()
    azioni["Ripristina righe rimosse"].trigger()
    assert m.valore(4, "Origine") == "IMPORTATO"


def test_filtri(finestra):
    pg = pagina(finestra, Scheda.CHECK_ACTUAL)
    pg.pannello.ricerca.setText("b1-bm")
    assert pg.pannello.proxy.rowCount() == 1 and pg.pannello.conteggio.text() == "1 di 5 righe"
    pg.pannello.ricerca.setText("")
    pg.pannello.seleziona_esito("Esclusa")
    assert pg.pannello.proxy.rowCount() == 1


def test_check_con_statistiche(finestra):
    finestra.progetto.imposta_check_sap(tabella(NEW[:4]))
    pg = pagina(finestra, Scheda.CHECK)
    assert pg.statistiche.isVisibleTo(finestra)
    assert pg.statistiche.voci["Mancante in SAP"].text() == "Mancante in SAP: 1  (20.0%)"
    pg.statistiche.voci["Mancante in SAP"].click()
    assert pg.pannello.proxy.rowCount() == 1


def test_esegui_con_attesa(app):
    assert esegui_con_attesa(None, "prova", lambda: 42) == 42
    with pytest.raises(ValueError, match="boom"):
        esegui_con_attesa(None, "prova", lambda: (_ for _ in ()).throw(ValueError("boom")))


def test_sap_script(finestra, tmp_path):
    pytest.importorskip("xlwt")
    from conversione_sedi.riconciliazione import Decisione
    p = finestra.progetto
    finestra.conferma = lambda testo: True
    finestra.informa = lambda titolo, testo: None
    pg = pagina(finestra, Scheda.SAP_SCRIPT)
    assert pg.avviso.text().startswith("Scheda bloccata") and not pg.b_genera.isEnabled()

    p.imposta_decisione("ITE-OT0B-1E-01-B1", Decisione.rimuovi())
    vuoto = pd.DataFrame(columns=["Sede tecnica", "Work Order", "Notification", "Descrizione", "Stato"])
    p.imposta_no_wo(vuoto, vuoto, p.sedi_per_no_wo())
    finestra.aggiorna()                                               # come dopo ogni azione nella GUI
    assert pg.b_genera.isEnabled()

    pg.b_genera.click()
    assert finestra.errori == []
    m = pg.pannello.modello
    assert [m.valore(r, "Tipo") for r in range(m.rowCount())] == ["Creazione", "Modifica", "Rimozione"]
    assert m.valore(1, "File") == "ZPM4R_FL_DES_MA_LOAD_mod_Step_2.xls"
    assert m.valore(0, "File") == "FL_da_creare_Step_1.xlsx"
    assert pg.dettaglio.modello.rowCount() == 1                       # dettaglio dello step 1 selezionato
    assert not pg.b_carica.isEnabled() and not pg.b_ok.isEnabled()
    assert pg.b_segna.text() == "Segna step 1 come caricato"

    pg.b_segna.click()
    assert m.valore(0, "Stato") == "Caricato - da validare" and pg.b_ok.isEnabled()
    assert m.stile(0, "Stato").sfondo is not None
    pg.b_ok.click()
    assert m.valore(0, "Stato") == "Validato" and pg.b_segna.text() == "Segna step 2 come caricato"

    # dati bloccati: la modifica in New viene rifiutata con un messaggio
    mn = pagina(finestra, Scheda.NEW).pannello.modello
    mn.setData(mn.index(2, 1), "Altra descrizione")
    assert "Interrompi piano" in str(finestra.errori[-1])

    pg = pagina(finestra, Scheda.SAP_SCRIPT)
    pg.b_interrompi.click()
    assert p.piano is None and pg.pannello.modello.rowCount() == 0


def test_sap_script_carica_in_sap(finestra, monkeypatch):
    pytest.importorskip("xlwt")
    from conversione_sedi import sap
    from conversione_sedi.riconciliazione import Decisione
    p = finestra.progetto
    finestra.conferma = lambda testo: True
    messaggi = []
    finestra.informa = lambda titolo, testo: messaggi.append(testo)
    caricati = []
    monkeypatch.setattr(sap, "connetti_sessione", lambda: "sessione")
    monkeypatch.setattr(sap, "carica_file", lambda s, tipo, file: caricati.append((tipo, file.name)) or "Elaborati 1 record")

    p.imposta_decisione("ITE-OT0B-1E-01-B1", Decisione.rimuovi())
    vuoto = pd.DataFrame(columns=["Sede tecnica", "Work Order", "Notification", "Descrizione", "Stato"])
    p.imposta_no_wo(vuoto, vuoto, p.sedi_per_no_wo())
    pg = pagina(finestra, Scheda.SAP_SCRIPT)
    pg.b_genera.click()
    assert not pg.b_carica.isEnabled()                               # step 1: creazione (programma esterno)
    assert "programma esterno" in pg.b_carica.toolTip()
    assert "Percorso troppo lungo" in pg.avviso.text()               # cartelle temporanee dei test > 128

    p.segna_caricato(1)
    p.registra_esito(1, True)
    finestra.aggiorna()
    assert pg.b_carica.isEnabled() and pg.b_carica.text() == "Carica step 2 in SAP"
    pg.b_carica.click()
    assert finestra.errori == []
    assert caricati == [(sap.TipoStep.MODIFICA, "ZPM4R_FL_DES_MA_LOAD_mod_Step_2.xls")]
    assert p.stati_step[1].value == "Caricato - da validare"
    assert "Elaborati 1 record" in messaggi[-1] and "Esito OK" in messaggi[-1]
