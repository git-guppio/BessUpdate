from pathlib import Path

import pytest

xlrd = pytest.importorskip("xlrd")
pytest.importorskip("xlwt")

from conversione_sedi.file_sap import (                     # noqa: E402
    country, intestazioni_cc04, intestazioni_zpm4r, nome_file, scrivi_piano,
)
from conversione_sedi.ingresso_uscita import ErroreDati     # noqa: E402
from conversione_sedi.piano_sap import Piano, Riga, Step, TipoStep as T   # noqa: E402

TEMPLATE = Path(__file__).resolve().parents[2] / "template_SAP"


def piano():
    return Piano(step=[
        Step(1, T.CREAZIONE, [Riga("USE-P1-01", "USE-P1-01", "Linea 01", "USE-P1")]),
        Step(2, T.MODIFICA, [Riga("USE-P1-TR-T1", "USE-P1-01-T1", "Trafo", "USE-P1-01", "USE-P1-TR-T1"),
                             Riga("USE-P1-01-T1-BM", "USE-P1-01-T1-R1", "Rack", "USE-P1-01-T1")]),
        Step(3, T.RIMOZIONE, [Riga("USE-P1-TR", "USA-P1-TR", "Vecchia", "USA", "USE-P1-TR")]),
        Step(4, T.DESCRIZIONE, [Riga("USE-P1-Z", "", "Descrizione è nuova", "", "USE-P1-Z")]),
    ])


def test_intestazioni_come_template():
    assert "\t".join(intestazioni_zpm4r()) == (TEMPLATE / "ZPM4R_CREA_FL_Template.txt").read_text(encoding="cp1252")
    assert len(intestazioni_zpm4r()) == 70 and "TEL_NUMBER_OLD " in intestazioni_zpm4r()
    sh = xlrd.open_workbook(str(TEMPLATE / "CC04_Template.xls")).sheet_by_index(0)
    assert intestazioni_cc04() == [sh.cell_value(0, c) for c in range(sh.ncols)]
    assert intestazioni_cc04()[-1] == "New Superior "


def test_nomi_e_country():
    p = piano()
    assert [nome_file(s) for s in p.step] == [
        "FL_da_creare_Step_1.xlsx", "ZPM4R_FL_DES_MA_LOAD_modificare_upload_Step_2.xls",
        "ZPM4R_FL_DES_MA_LOAD_rimuovere_upload_Step_3.xls", "ZPM4R_CREA_FL_modifica_upload_Step_4.txt"]
    assert country("USE-US64-1A") == "US"


def test_scrivi_piano(tmp_path):
    cartella = tmp_path / "SAP_USE-P1"
    scritti = scrivi_piano(piano(), cartella)
    assert sorted(p.name for p in cartella.iterdir()) == [
        "FL_da_creare_Step_1.xlsx", "ZPM4R_CREA_FL_modifica_upload_Step_4.txt",
        "ZPM4R_FL_DES_MA_LOAD_modificare_upload_Step_2.xls", "ZPM4R_FL_DES_MA_LOAD_rimuovere_upload_Step_3.xls"]

    # creazione: file Excel per il programma esterno, colonne FL e Descriptions
    from openpyxl import load_workbook
    ws = load_workbook(scritti[1]).active
    assert [[c.value for c in r] for r in ws.iter_rows()] == [["FL", "Descriptions"], ["USE-P1-01", "Linea 01"]]
    assert ws["A2"].number_format == "@"

    sh = xlrd.open_workbook(str(scritti[2])).sheet_by_index(0)
    assert sh.name == "Foglio1"
    assert [sh.row_values(r) for r in range(sh.nrows)] == [
        ["Functional Location", "New Code", "Description (if required)", "New Superior "],
        ["USE-P1-TR-T1", "USE-P1-01-T1", "Trafo", "USE-P1-01"],
        ["USE-P1-01-T1-BM", "USE-P1-01-T1-R1", "Rack", "USE-P1-01-T1"]]
    sh = xlrd.open_workbook(str(scritti[3])).sheet_by_index(0)
    assert sh.row_values(1) == ["USE-P1-TR", "USA-P1-TR", "Vecchia", "USA"]

    dati = scritti[4].read_bytes()
    assert not dati.startswith(b"\xef\xbb\xbf")                      # senza BOM
    righe = dati.decode("cp1252").split("\r\n")
    assert righe[0] == (TEMPLATE / "ZPM4R_CREA_FL_Template.txt").read_text(encoding="cp1252")
    valori = righe[1].split("\t")
    colonne = intestazioni_zpm4r()
    assert len(valori) == 70
    assert {colonne[i]: v for i, v in enumerate(valori) if v} == {
        "TPLNR": "USE-P1-Z", "PLTXT": "Descrizione è nuova", "COUNTRY": "US"}
    assert righe[-1] == ""                                           # CRLF finale


def test_sovrascrive_generazione_precedente(tmp_path):
    vecchio = tmp_path / "ZPM4R_FL_DES_MA_LOAD_modificare_upload_Step_9.xls"
    vecchio.write_bytes(b"x")
    vecchio_nome = tmp_path / "CC04_modificare_upload_Step_2.xls"   # nome delle versioni precedenti
    vecchio_nome.write_bytes(b"x")
    altro = tmp_path / "mio_file.xls"
    altro.write_bytes(b"x")
    scrivi_piano(piano(), tmp_path)
    assert not vecchio.exists() and not vecchio_nome.exists() and altro.exists()


def test_caratteri_non_ammessi(tmp_path):
    p = Piano(step=[Step(1, T.DESCRIZIONE, [Riga("USE-P1-Z", "", "Più → meno", "", "")])])
    with pytest.raises(ErroreDati, match="Windows-1252: USE-P1-Z"):
        scrivi_piano(p, tmp_path)


def test_piano_con_errori(tmp_path):
    with pytest.raises(ErroreDati, match="file non generati"):
        scrivi_piano(Piano(errori=["x"]), tmp_path)
