"""Esportazione di una tabella in Excel con i colori della GUI.

Tutte le celle sono in formato Testo (valori come "++E-..." o "01" non
vengono interpretati come formule o numeri).
"""
from __future__ import annotations

from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from conversione_sedi.gui.stile import Stile
from conversione_sedi.ingresso_uscita import ErroreDati

LARGHEZZA_MAX = 60


def _rgb(colore: str) -> str:
    return "FF" + colore.lstrip("#").upper()


def esporta_excel(df: pd.DataFrame, percorso: str | Path, foglio: str = "Dati",
                  stile: Callable[[int, str], Stile] | None = None) -> Path:
    """Scrive df in un file .xlsx; `stile(riga, colonna)` fornisce i colori."""
    percorso = Path(percorso)
    if percorso.suffix.lower() != ".xlsx":
        percorso = percorso.with_name(percorso.name + ".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = foglio[:31]

    colonne = [str(c) for c in df.columns]
    ws.append(colonne)
    for cella in ws[1]:
        cella.font = Font(bold=True)
        cella.number_format = "@"

    larghezze = [len(c) for c in colonne]
    for i, riga in enumerate(df.itertuples(index=False)):
        for j, valore in enumerate(riga):
            testo = "" if valore is None or (isinstance(valore, float) and pd.isna(valore)) else str(valore)
            cella = ws.cell(row=i + 2, column=j + 1, value=testo)
            cella.data_type = "s"                   # mai formula
            cella.number_format = "@"
            larghezze[j] = max(larghezze[j], len(testo))
            if stile:
                s = stile(i, colonne[j])
                if s.testo or s.barrato:
                    cella.font = Font(color=_rgb(s.testo) if s.testo else None, strike=s.barrato)
                if s.sfondo:
                    cella.fill = PatternFill("solid", start_color=_rgb(s.sfondo))

    for j, larghezza in enumerate(larghezze):
        ws.column_dimensions[ws.cell(row=1, column=j + 1).column_letter].width = min(larghezza + 2, LARGHEZZA_MAX)
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(percorso)
    except PermissionError:
        raise ErroreDati([f"Impossibile scrivere {percorso.name}: il file è aperto in un altro programma?"]) from None
    return percorso


def testo_tabulato(df: pd.DataFrame, intestazioni: bool = True) -> str:
    """Testo separato da tabulazioni, incollabile in Excel (per gli appunti)."""
    righe = ["\t".join(map(str, df.columns))] if intestazioni else []
    for riga in df.itertuples(index=False):
        righe.append("\t".join("" if v is None or (isinstance(v, float) and pd.isna(v))
                               else str(v).replace("\t", " ").replace("\n", " ") for v in riga))
    return "\r\n".join(righe)
